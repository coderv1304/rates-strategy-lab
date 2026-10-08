"""Create excel/RatesDashboard.xlsx (Data, Hedge, Curve sheets). Add the VBA module afterwards.

Run:  python -m rateslab.build_excel
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from . import config as C


def build() -> str:
    yields = pd.read_csv(C.DATA_PROC / "yields.csv", parse_dates=["Date"])
    wb = Workbook()

    data = wb.active
    data.title = "Data"
    data.append(list(yields.columns))
    for rec in yields.itertuples(index=False):
        data.append([rec[0].to_pydatetime()] + [float(v) for v in rec[1:]])
    for row in data.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "yyyy-mm-dd"
    for c in data[1]:
        c.font = Font(bold=True)
    data.freeze_panes = "B2"

    hedge = wb.create_sheet("Hedge")
    inputs = [("Valuation date", None), ("2Y par yield (%)", 4.00), ("10Y par yield (%)", 4.40),
              ("2Y notional (USD)", 100_000_000)]
    hedge["A1"], hedge["B1"] = "INPUTS (blue cells)", ""
    for i, (label, val) in enumerate(inputs, start=2):
        hedge.cell(i, 1, label)
        cell = hedge.cell(i, 2, val)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    hedge["B2"] = pd.Timestamp(yields["Date"].iloc[-1]).to_pydatetime()
    hedge["B2"].number_format = "yyyy-mm-dd"
    hedge["B5"].number_format = "#,##0"
    hedge["A7"] = "OUTPUTS (filled by the ComputeHedge macro)"
    for r, label in enumerate(["2Y modified duration", "10Y modified duration", "2Y DV01 per 100 face",
                               "10Y DV01 per 100 face", "10Y notional (USD) for DV01 neutrality",
                               "Hedge ratio (10Y per 2Y notional)", "Net DV01 (USD per bp)"], start=8):
        hedge.cell(r, 1, label)
    for cell in (hedge["A1"], hedge["A7"]):
        cell.font = Font(bold=True)
    hedge.column_dimensions["A"].width = 42
    hedge.column_dimensions["B"].width = 20
    hedge["A16"] = (
        "Buttons: Developer > Insert > Button, then assign"
        " RefreshData / FillFromLatest / ComputeHedge / ChartCurve."
    )
    hedge["A16"].alignment = Alignment(wrap_text=True)

    wb.create_sheet("Curve")
    out = C.ROOT / "excel" / "RatesDashboard.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    return str(out)


if __name__ == "__main__":
    print("Wrote", build())
